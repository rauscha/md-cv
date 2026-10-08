import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document

import datetime as dt

from build_cv import (
    parse_cv, render_docx, BULLET_PREFIX, HEAD_INDENT, INDENT, NUM_HANG, NUM_INDENT,
    RIGHT_MARGIN, SIDE_MARGIN, SPACE_GAP, TITLE,
)

MD = """# Jane Doe, MD

Email: jane@example.edu

## APPOINTMENTS

2021-present | Assistant Professor, **Some University**

## PUBLICATIONS

### (a) Peer-reviewed Publications

1. Doe J. A paper. *Journal*: 1-2. 2024

### (b) Book Chapters
"""


def build(tmp_path):
    out = tmp_path / "cv.docx"
    render_docx(parse_cv(MD), out)
    return Document(str(out))


def texts(doc):
    return [p.text for p in doc.paragraphs]


def test_structure_in_order(tmp_path):
    t = texts(build(tmp_path))
    assert t[0] == TITLE
    assert t[1] == "Jane Doe, MD"
    assert "APPOINTMENTS" in t
    assert "2021-present\tAssistant Professor, Some University" in t
    assert "(a) Peer-reviewed Publications" in t


def test_dated_entry_hanging_indent(tmp_path):
    doc = build(tmp_path)
    p = next(p for p in doc.paragraphs if p.text.startswith("2021-present\t"))
    pf = p.paragraph_format
    assert pf.left_indent == INDENT
    assert pf.first_line_indent == -INDENT
    assert pf.tab_stops[0].position == INDENT


def test_bold_italic_runs(tmp_path):
    doc = build(tmp_path)
    entry = next(p for p in doc.paragraphs if "Some University" in p.text)
    assert any(r.bold and r.text == "Some University" for r in entry.runs)
    cite = next(p for p in doc.paragraphs if p.text.startswith("1.\tDoe J"))
    assert any(r.italic and r.text == "Journal" for r in cite.runs)


def test_empty_subsection_renders_none(tmp_path):
    t = texts(build(tmp_path))
    i = t.index("(b) Book Chapters")
    assert t[i + 1] == "None"


def test_title_block_is_centered_bold(tmp_path):
    doc = build(tmp_path)
    for p in doc.paragraphs[:2]:
        assert p.alignment == 1  # WD_ALIGN_PARAGRAPH.CENTER
        assert p.runs[0].bold


def test_margins(tmp_path):
    doc = build(tmp_path)
    section = doc.sections[0]
    assert section.left_margin == SIDE_MARGIN
    assert section.right_margin == RIGHT_MARGIN


def _bullet_doc(tmp_path, md):
    out = tmp_path / "bullets.docx"
    render_docx(parse_cv(md), out)
    return Document(str(out))


def test_sub_bullet_shares_the_description_column(tmp_path):
    doc = _bullet_doc(
        tmp_path,
        "## TEACHING ACTIVITIES\nResidency, Some University\n- Designed a curriculum\n",
    )
    p = next(p for p in doc.paragraphs if "Designed a curriculum" in p.text)
    pf = p.paragraph_format
    assert pf.left_indent == INDENT
    assert pf.first_line_indent is None
    assert len(pf.tab_stops) == 0
    assert p.text == BULLET_PREFIX + "Designed a curriculum"


def test_sub_bullet_honors_inline_markup(tmp_path):
    doc = _bullet_doc(tmp_path, "## TEACHING\n- Ran the **FUNGI** simulation\n")
    p = next(p for p in doc.paragraphs if "FUNGI" in p.text)
    assert any(r.bold and r.text == "FUNGI" for r in p.runs)


def test_sub_bullet_gap_becomes_space_after(tmp_path):
    doc = _bullet_doc(tmp_path, "## TEACHING\n- First bullet\n\nNext block\n")
    p = next(p for p in doc.paragraphs if "First bullet" in p.text)
    assert p.paragraph_format.space_after == SPACE_GAP


def test_gap_becomes_space_after_on_plain_paragraph(tmp_path):
    out = tmp_path / "gap.docx"
    render_docx(parse_cv("## PUBS\n\nDoe J. One. 2024\n\nDoe J. Two. 2023\n"), out)
    doc = Document(str(out))
    gapped = next(p for p in doc.paragraphs if p.text == "Doe J. One. 2024")
    assert gapped.paragraph_format.space_after == SPACE_GAP


def test_scaffold_section_not_rendered(tmp_path):
    out = tmp_path / "s.docx"
    render_docx(parse_cv("## REAL\n\n2020 | A\n\n## GRANTS [TBC]\n\n2021 | Grant\n"), out)
    t = " ".join(p.text for p in Document(str(out)).paragraphs)
    assert "GRANTS" not in t and "Grant" not in t
    assert "2020" in t


def test_section_titles_are_uppercased_and_subsections_italic_underlined(tmp_path):
    out = tmp_path / "c.docx"
    render_docx(parse_cv("## Bibliography\n\n### (a) Papers\n\n1. A paper\n"), out)
    t = texts(Document(str(out)))
    assert "BIBLIOGRAPHY" in t
    sub = next(p for p in Document(str(out)).paragraphs if p.text == "(a) Papers")
    assert all(r.italic and r.underline for r in sub.runs)


def test_numbered_item_hangs_its_number(tmp_path):
    doc = build(tmp_path)
    pf = next(p for p in doc.paragraphs if p.text.startswith("1.\t")).paragraph_format
    assert pf.left_indent == NUM_INDENT
    assert pf.first_line_indent == -NUM_HANG


def test_contact_labels_and_dated_line(tmp_path):
    out = tmp_path / "h.docx"
    md = "# Jane Doe\n\nDated: Chicago\nAddress: One St\nChicago, IL\nEmail: j@x.edu\n"
    render_docx(parse_cv(md), out, cv_date=dt.date(2014, 4, 6))
    doc = Document(str(out))
    t = texts(doc)
    assert "Chicago, 4/6/14" in t
    assert "Address:\tOne St" in t and "Email:\tj@x.edu" in t
    cont = next(p for p in doc.paragraphs if p.text == "Chicago, IL")
    assert cont.paragraph_format.left_indent == INDENT


def test_footer_carries_name_month_and_page_fields(tmp_path):
    out = tmp_path / "f.docx"
    render_docx(parse_cv("# Jane Doe, MD\n\n## A\n\nx\n"), out, cv_date=dt.date(2014, 4, 6))
    footer = Document(str(out)).sections[0].footer
    xml = footer._element.xml
    assert "Jane Doe MD\tApril 2014\tPage " in footer.paragraphs[0].text
    assert "PAGE" in xml and "NUMPAGES" in xml


def test_heading_line_right_aligns_dates_and_indents_its_bullets(tmp_path):
    out = tmp_path / "g.docx"
    render_docx(parse_cv("## FUNDING\n\nNIH K12 | 2009-2016\n- Role: PI\n"), out)
    doc = Document(str(out))
    head = next(p for p in doc.paragraphs if p.text.startswith("NIH K12"))
    assert head.text == "NIH K12\t2009-2016"
    assert head.runs[0].bold
    bullet = next(p for p in doc.paragraphs if p.text == "Role: PI")
    assert bullet.paragraph_format.left_indent == HEAD_INDENT


def test_page_break_lands_on_next_heading(tmp_path):
    out = tmp_path / "p.docx"
    render_docx(parse_cv("## A\n\nx\n\n---\n## STATEMENT\n\ny\n"), out)
    p = next(p for p in Document(str(out)).paragraphs if p.text == "STATEMENT")
    assert p.paragraph_format.page_break_before
