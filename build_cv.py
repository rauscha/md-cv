#!/usr/bin/env python3
"""Build a styled .docx and .pdf academic CV from a Markdown source.

Canonical home: https://github.com/rauscha/md-cv
Usage: python build_cv.py path/to/CV.md [-o OUTDIR] [--date YYYY-MM-DD]

Layout follows the University of Chicago BSD CV format (the COAP template):
centered CURRICULUM VITAE title, bold all-caps section headings,
italic-underlined lettered subsections, dates in a 1" left column, numbered
citations, and a "Name / Month Year / Page X of Y" footer.

Format: `# Name` then contact lines (`Label: value` puts the label in its own
column; unlabeled lines after a labeled one continue its value; `Dated: Place`
prints "Place, M/D/YY" right-aligned under the name, using the build date);
`## SECTION`; `### (a) Subsection`; dated entries as `2018-2021 | description`
(must start with a 4-digit year, or the word `Current` for an ongoing role);
`Title | 2009-2016` (text first, dates last) is a heading line with the dates
right-aligned, as for a grant; `1. text` is a numbered item with a hanging
indent; `- text` is a sub-bullet indented into the text column of the line that
introduces it; an indented `  - text` is a nested item with a bullet glyph;
`---` on its own line starts the next item on a new page; a blank line before an
entry adds extra vertical space after the preceding item (a gap); everything
else is a plain paragraph. `**bold**`, `*italic*`, `***bold italic***` and
`<u>underline</u>` inline.
`[TBC]` on a normal line strips just that line from the output. `[TBC]` on a
`##`/`###` header instead marks the whole section/subsection as a scaffold:
the header and everything under it stay in the source but are entirely
suppressed from the rendered output until `[TBC]` is removed from the header.
"""
from __future__ import annotations

import argparse
import datetime as dt
import difflib
import re
import shutil
import subprocess
import sys
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_UNDERLINE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt, Twips

# ---------------------------------------------------------------------------
# STYLE CONSTANTS — the entire visual design lives here. Matches the UChicago
# BSD / COAP CV template structure (single spacing, 1" date column) in our own
# typeface.
# ---------------------------------------------------------------------------
FONT_NAME = "Calibri"
FOOTER_FONT_NAME = FONT_NAME
BODY_SIZE = Pt(11)
NAME_SIZE = Pt(11)          # the title block is bold body-size text, not a large name
TITLE = "CURRICULUM VITAE"
INDENT = Inches(1.0)        # description column: hanging indent + tab stop for dated
                            # entries and contact labels, left indent for sub-bullets
NUM_INDENT = Twips(450)    # numbered items: text column
NUM_HANG = Inches(0.25)     # numbered items: the number hangs this far left of the text
HEAD_INDENT = Inches(0.25)  # `Title | dates` heading lines (grants) and their sub-bullets
GLYPH_HANG = Inches(0.25)   # nested `  - ` items: glyph hangs this far left of the text
BULLET_PREFIX = ""          # sub-bullets sit in the text column with no glyph
NESTED_GLYPH = "•"
SIDE_MARGIN = Inches(1.0)   # left margin
RIGHT_MARGIN = Twips(990)  # the template runs its right margin narrower than the left
TOP_BOTTOM_MARGIN = Inches(1.0)
SPACE_BEFORE_PARA = Pt(0)
SPACE_BEFORE_SECTION = Pt(12)   # one blank line above each section heading
SPACE_AFTER_SECTION = Pt(0)
SPACE_BEFORE_SUBSECTION = Pt(8)
SPACE_AFTER_PARA = Pt(0)
SPACE_GAP = Pt(11)          # extra space_after on a paragraph followed by a blank line in the source


@dataclass
class Entry:
    dates: str
    text: str
    gap: bool = False


@dataclass
class Para:
    text: str
    gap: bool = False


@dataclass
class Bullet:
    """A `- ` sub-bullet: indented into the text column, no hanging indent."""

    text: str
    gap: bool = False


@dataclass
class Nested:
    """An indented `  - ` item: a glyph bullet one step inside the text column."""

    text: str
    gap: bool = False


@dataclass
class Numbered:
    """A `1. ` item: the number hangs left of a wrapped text block."""

    number: str
    text: str
    gap: bool = False


@dataclass
class Heading:
    """A `Title | 2009-2016` line: bold title with the dates right-aligned."""

    text: str
    dates: str
    gap: bool = False


@dataclass
class PageBreak:
    """`---`: the next rendered paragraph starts on a new page."""

    gap: bool = False


@dataclass
class Subsection:
    title: str
    items: list = field(default_factory=list)
    scaffold: bool = False


@dataclass
class Section:
    title: str
    items: list = field(default_factory=list)
    scaffold: bool = False


@dataclass
class CV:
    name: str
    contact: list
    sections: list


DATES = r"(?:\d{4}|Current)[^|]*"
DATE_ENTRY_RE = re.compile(rf"^({DATES})\|(.*)$")
HEADING_DATES = r"(?:\d{4}|Current)(?:[\s\-\u2013\u2014,]*(?:\d{4}|present|Present))*[\s\-\u2013\u2014]*"
HEADING_RE = re.compile(rf"^([^|]+?)\s*\|\s*({HEADING_DATES})$")
NUMBERED_RE = re.compile(r"^(\d{1,3})\.\s+(.*)$")
PAGE_BREAK_RE = re.compile(r"^-{3,}$")
LABEL_RE = re.compile(r"^([A-Za-z][A-Za-z ]{0,19}):\s+(.*)$")


def parse_cv(text: str) -> CV:
    name = ""
    contact: list[str] = []
    sections: list[Section] = []
    section: Section | None = None
    sub: Subsection | None = None
    saw_blank = False
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            saw_blank = True
            continue
        is_header = line.startswith("### ") or line.startswith("## ") or line.startswith("# ")
        if "[TBC]" in line and not is_header:
            # Non-header [TBC] lines are just dropped from the output.
            continue
        if line.startswith("### "):
            if section is None:
                raise ValueError(
                    f"subsection before any '## SECTION' header: {line!r}"
                )
            sub = Subsection(line[4:].strip(), scaffold="[TBC]" in line)
            section.items.append(sub)
            saw_blank = False
        elif line.startswith("## "):
            section = Section(line[3:].strip(), scaffold="[TBC]" in line)
            sub = None
            sections.append(section)
            saw_blank = False
        elif line.startswith("# "):
            name = line[2:].strip()
            saw_blank = False
        else:
            if sub is not None:
                target = sub.items
            elif section is not None:
                target = section.items
            else:
                contact.append(line)
                saw_blank = False
                continue
            m = DATE_ENTRY_RE.match(line)
            h = HEADING_RE.match(line)
            n = NUMBERED_RE.match(line)
            if m:
                item = Entry(m.group(1).strip(), m.group(2).strip())
            elif h:
                item = Heading(h.group(1).strip(), h.group(2).strip())
            elif PAGE_BREAK_RE.match(line):
                item = PageBreak()
            elif line.startswith("- ") and raw[:1] in (" ", "\t"):
                item = Nested(line[2:].strip())
            elif line.startswith("- "):
                item = Bullet(line[2:].strip())
            elif n:
                item = Numbered(n.group(1), n.group(2).strip())
            else:
                item = Para(line)
            if saw_blank and target:
                target[-1].gap = True
            target.append(item)
            saw_blank = False
    return CV(name, contact, sections)


# Lookarounds enforce CommonMark-style flanking (no adjacent word char or `*`
# just outside the delimiters), so co-author asterisks like `Doe J*,` stay literal.
TOKEN_RE = re.compile(
    r"(<u>.+?</u>"
    r"|\*\*\*[^*]+\*\*\*"
    r"|\*\*[^*]+\*\*"
    r"|(?<![\w*])\*[^*\s][^*]*\*(?![\w*]))"
)


def tokenize_runs(text: str) -> list[tuple[str, bool, bool, bool]]:
    """Split inline markup into (text, bold, italic, underline) runs.

    `<u>` may wrap bold/italic markup (`<u>**Doe J**</u>`); the reverse nesting
    is not supported.
    """
    out = []
    for part in TOKEN_RE.split(text):
        if not part:
            continue
        if part.startswith("<u>") and part.endswith("</u>"):
            out.extend((t, b, i, True) for t, b, i, _ in tokenize_runs(part[3:-4]))
        elif part.startswith("***") and part.endswith("***") and len(part) > 6:
            out.append((part[3:-3], True, True, False))
        elif part.startswith("**") and part.endswith("**"):
            out.append((part[2:-2], True, False, False))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            out.append((part[1:-1], False, True, False))
        else:
            out.append((part, False, False, False))
    return out


def _add_runs(p, text: str, bold: bool = False):
    for t, b, italic, underline in tokenize_runs(text):
        r = p.add_run(t)
        r.bold = b or bold
        r.italic = italic
        if underline:
            r.underline = WD_UNDERLINE.SINGLE
    return p


def _hanging(p, indent, hang=None, tab=None):
    pf = p.paragraph_format
    pf.left_indent = indent
    pf.first_line_indent = -(hang if hang is not None else indent)
    pf.tab_stops.add_tab_stop(tab if tab is not None else indent, WD_TAB_ALIGNMENT.LEFT)
    return pf


class _Renderer:
    def __init__(self, doc, text_width):
        self.doc = doc
        self.text_width = text_width
        self.break_next = False
        self.bullet_indent = INDENT

    def para(self):
        p = self.doc.add_paragraph()
        if self.break_next:
            p.paragraph_format.page_break_before = True
            self.break_next = False
        return p

    def items(self, items):
        for item in items:
            if isinstance(item, Subsection):
                if item.scaffold:
                    continue
                p = self.para()
                p.paragraph_format.space_before = SPACE_BEFORE_SUBSECTION
                p.paragraph_format.keep_with_next = True
                _add_runs(p, item.title)
                for r in p.runs:
                    r.italic = True
                    r.underline = WD_UNDERLINE.SINGLE
                self.bullet_indent = INDENT
                if item.items:
                    self.items(item.items)
                else:
                    self.para().add_run("None")
                continue
            if isinstance(item, PageBreak):
                self.break_next = True
                continue
            p = self.para()
            pf = p.paragraph_format
            if isinstance(item, Entry):
                _hanging(p, INDENT)
                p.add_run(item.dates + "\t")
                _add_runs(p, item.text)
                self.bullet_indent = INDENT
            elif isinstance(item, Heading):
                pf.left_indent = HEAD_INDENT
                pf.keep_with_next = True
                pf.tab_stops.add_tab_stop(Emu(self.text_width - HEAD_INDENT), WD_TAB_ALIGNMENT.RIGHT)
                _add_runs(p, item.text, bold=True)
                p.add_run("\t" + item.dates)
                self.bullet_indent = HEAD_INDENT
            elif isinstance(item, Bullet):
                pf.left_indent = self.bullet_indent
                _add_runs(p, BULLET_PREFIX + item.text)
            elif isinstance(item, Nested):
                _hanging(p, INDENT + GLYPH_HANG, GLYPH_HANG)
                p.add_run(NESTED_GLYPH + "\t")
                _add_runs(p, item.text)
            elif isinstance(item, Numbered):
                _hanging(p, NUM_INDENT, NUM_HANG)
                p.add_run(item.number + ".\t")
                _add_runs(p, item.text)
                self.bullet_indent = NUM_INDENT
            else:
                _add_runs(p, item.text)
                self.bullet_indent = INDENT
            if item.gap:
                pf.space_after = SPACE_GAP


def _add_field(p, instr: str, bold: bool = False):
    """Append a Word field (PAGE, NUMPAGES) as a complex field run sequence."""
    for kind in ("begin", None, "separate", "1", "end"):
        r = p.add_run()
        r.bold = bold
        if kind is None:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {instr} "
        elif kind == "1":
            r.text = "1"
            continue
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        r._r.append(el)


def _footer(section, name: str, month_year: str, text_width):
    p = section.footer.paragraphs[0]
    pf = p.paragraph_format
    pf.tab_stops.add_tab_stop(Emu(text_width // 2), WD_TAB_ALIGNMENT.CENTER)
    pf.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
    p.add_run(f"{name}\t{month_year}\tPage ")
    _add_field(p, "PAGE", bold=True)
    p.add_run(" of ")
    _add_field(p, "NUMPAGES", bold=True)
    for r in p.runs:
        r.font.name = FOOTER_FONT_NAME
        r.font.size = BODY_SIZE


def _contact(doc, lines, cv_date: dt.date):
    labeled = False
    for line in lines:
        if line.startswith("Dated:"):
            continue
        m = LABEL_RE.match(line)
        p = doc.add_paragraph()
        if m:
            labeled = True
            _hanging(p, INDENT)
            p.add_run(m.group(1) + ":\t")
            _add_runs(p, m.group(2))
        else:
            if labeled:
                p.paragraph_format.left_indent = INDENT
            _add_runs(p, line)


def render_docx(cv: CV, path: Path, cv_date: dt.date | None = None) -> None:
    cv_date = cv_date or dt.date.today()
    doc = Document()
    text_width = None
    for section in doc.sections:
        section.left_margin = SIDE_MARGIN
        section.right_margin = RIGHT_MARGIN
        section.top_margin = TOP_BOTTOM_MARGIN
        section.bottom_margin = TOP_BOTTOM_MARGIN
        text_width = section.page_width - SIDE_MARGIN - RIGHT_MARGIN
        _footer(section, cv.name.replace(",", ""), cv_date.strftime("%B %Y"), text_width)
    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = BODY_SIZE
    normal.paragraph_format.space_before = SPACE_BEFORE_PARA
    normal.paragraph_format.space_after = SPACE_AFTER_PARA
    normal.paragraph_format.line_spacing = 1.0

    for text in (TITLE, cv.name):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold = True
        r.font.size = NAME_SIZE

    place = next((c[len("Dated:"):].strip() for c in cv.contact if c.startswith("Dated:")), None)
    if place is not None:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        stamp = f"{cv_date.month}/{cv_date.day}/{cv_date:%y}"
        p.add_run(f"{place}, {stamp}" if place else stamp)

    if cv.contact:
        doc.add_paragraph()
        _contact(doc, cv.contact, cv_date)

    renderer = _Renderer(doc, text_width)
    for section in cv.sections:
        if section.scaffold:
            continue
        p = renderer.para()
        p.paragraph_format.space_before = SPACE_BEFORE_SECTION
        p.paragraph_format.space_after = SPACE_AFTER_SECTION
        p.paragraph_format.keep_with_next = True
        p.add_run(section.title.upper()).bold = True
        renderer.bullet_indent = INDENT
        renderer.items(section.items)

    doc.save(str(path))


def _export_via_word(docx_path: Path, pdf_path: Path) -> bool:
    try:
        import win32com.client
    except ImportError:
        return False
    try:
        word = win32com.client.DispatchEx("Word.Application")
    except Exception:
        return False
    try:
        word.Visible = False
        doc = word.Documents.Open(str(docx_path.resolve()))
        try:
            doc.SaveAs2(str(pdf_path.resolve()), FileFormat=17)  # wdFormatPDF
        finally:
            doc.Close(False)
            del doc  # release proxy while Word is still alive (avoids RPC warnings)
    except Exception as exc:
        print(f"Word export failed ({exc}); trying LibreOffice...", file=sys.stderr)
        return False
    finally:
        word.Quit()
        del word
    return pdf_path.exists()


def _export_via_soffice(docx_path: Path, pdf_path: Path, profile_dir: Path | None = None) -> bool:
    soffice = shutil.which("soffice")
    if not soffice:
        for candidate in (
            Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
            Path("/Applications/LibreOffice.app/Contents/MacOS/soffice"),
            Path("/usr/bin/soffice"),
        ):
            if candidate.exists():
                soffice = str(candidate)
                break
    if not soffice:
        return False
    cmd = [soffice]
    if profile_dir is not None:
        # Per-request profile dir avoids LibreOffice's single-profile lock when
        # multiple requests run headless soffice concurrently (matches
        # converter/server.js's `-env:UserInstallation=file://{profile}` flag).
        cmd.append(f"-env:UserInstallation=file://{profile_dir}")
    cmd += ["--headless", "--convert-to", "pdf",
            "--outdir", str(pdf_path.parent), str(docx_path)]
    try:
        subprocess.run(cmd, check=True, timeout=120)
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as exc:
        print(f"LibreOffice export failed ({exc})", file=sys.stderr)
        return False
    produced = pdf_path.parent / (docx_path.stem + ".pdf")
    if produced != pdf_path and produced.exists():
        produced.replace(pdf_path)
    if not pdf_path.exists():
        # soffice exits 0 even when it cannot load the file (e.g. Writer not installed)
        print("LibreOffice export failed (no PDF was written)", file=sys.stderr)
        return False
    return True


def export_pdf(docx_path: Path, pdf_path: Path, profile_dir: Path | None = None) -> str:
    # Delete the previous PDF first: each exporter is judged by whether the PDF exists
    # afterward, so a stale copy left in place would pass a failed export as a success.
    try:
        pdf_path.unlink(missing_ok=True)
    except PermissionError:
        sys.exit(f"ERROR: cannot replace {pdf_path.name}; close it in any PDF viewer and rerun.")
    if _export_via_word(docx_path, pdf_path):
        return "word"
    if _export_via_soffice(docx_path, pdf_path, profile_dir=profile_dir):
        return "libreoffice"
    raise RuntimeError(
        "PDF export failed: no exporter produced a PDF. Install Microsoft Word (plus "
        "`pip install pywin32`) or LibreOffice (including its Writer component)."
    )


def extract_pdf_text(pdf_path: Path):
    from pypdf import PdfReader

    reader = PdfReader(str(pdf_path))
    text = "\n".join((page.extract_text() or "") for page in reader.pages)
    return text, len(reader.pages)


class PlaceholderError(Exception):
    """Raised when the built PDF still contains unresolved placeholder text.

    A normal thing for a colleague filling in the sample CV to trigger — callers
    (e.g. a web service) should treat this as a client-error condition, not an
    internal failure.
    """


def check_for_placeholders(text: str) -> None:
    """Raise PlaceholderError if built output text still contains a placeholder.

    Call this wherever the built PDF's text is available (see `main`, which
    calls it right after `extract_pdf_text`) so a library caller gets a
    catchable exception instead of the process exiting under it.
    """
    if "TBC" in text:
        raise PlaceholderError("[TBC] content leaked into the PDF")


def main(argv=None):
    ap = argparse.ArgumentParser(description="Build a .docx and .pdf CV from Markdown.")
    ap.add_argument("markdown", type=Path)
    ap.add_argument("-o", "--outdir", type=Path, default=None)
    ap.add_argument(
        "--profile-dir", type=Path, default=None,
        help="Per-request LibreOffice profile dir, passed through to the soffice "
             "exporter as -env:UserInstallation (see _export_via_soffice). Lets a "
             "server give each concurrent request its own profile instead of sharing "
             "the default one, which headless soffice corrupts under concurrency. "
             "Default: none (unchanged CLI/Word behaviour).",
    )
    ap.add_argument(
        "--date", type=dt.date.fromisoformat, default=None,
        help="Date printed on the CV (YYYY-MM-DD): the right-aligned `Dated:` line "
             "and the footer's month. Default: today.",
    )
    args = ap.parse_args(argv)

    outdir = args.outdir or args.markdown.parent
    outdir.mkdir(parents=True, exist_ok=True)
    docx_path = outdir / (args.markdown.stem + ".docx")
    pdf_path = outdir / (args.markdown.stem + ".pdf")

    old_text = extract_pdf_text(pdf_path)[0] if pdf_path.exists() else None

    cv = parse_cv(args.markdown.read_text(encoding="utf-8"))
    render_docx(cv, docx_path, cv_date=args.date)
    engine = export_pdf(docx_path, pdf_path, profile_dir=args.profile_dir)

    new_text, pages = extract_pdf_text(pdf_path)
    try:
        check_for_placeholders(new_text)
    except PlaceholderError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        # Exit 3, not 2: argparse's own ArgumentParser.error() calls self.exit(2, ...)
        # for CLI usage errors, so 2 is already spoken for. Reusing it here would make
        # a malformed invocation indistinguishable from a leaked placeholder to any
        # caller (e.g. converter/server.js) that maps exit codes to HTTP responses.
        sys.exit(3)
    print(f"Built {docx_path.name} and {pdf_path.name} via {engine}; {pages} page(s).")
    if old_text is not None:
        diff = list(difflib.unified_diff(
            old_text.splitlines(), new_text.splitlines(),
            "previous.pdf", "new.pdf", lineterm=""))
        print("\n".join(diff) if diff else "No text changes vs previous PDF.")


if __name__ == "__main__":
    main()
